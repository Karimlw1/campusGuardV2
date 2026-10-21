"""
Features: 2-Lockdown, 3-Photo, 4-QR compromise, 5-Guard stats,
          7-Visitor pre-reg, 8-Bulk import, 9-Audit log, 12-QR expiry
"""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, BackgroundTasks
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc, func
from datetime import datetime, timedelta, date
from typing import Optional, List
from pydantic import BaseModel
import secrets, json, io, os, shutil
from app.db.session import get_db
from app.models.models import User, Gate, Alert, Entry, AuditLog, SystemState, Visitor
from app.core.dependencies import require_admin, require_guard, get_current_user
from app.core.security import hash_password
from app.utils.qr import generate_qr_hash, generate_qr_image

router = APIRouter(tags=["features"])

# ── Audit helper ──────────────────────────────────────────────────────────────
def audit(db, actor, action, target_type=None, target_id=None, detail=None):
    log = AuditLog(
        actor_id=actor.id if actor else None,
        actor_name=actor.full_name if actor else "System",
        action=action,
        target_type=target_type,
        target_id=target_id,
        detail=json.dumps(detail) if detail else None,
    )
    db.add(log)

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 2 — EMERGENCY LOCKDOWN
# ══════════════════════════════════════════════════════════════════════════════
class LockdownRequest(BaseModel):
    active: bool
    reason: Optional[str] = "Emergency lockdown"

@router.get("/lockdown/status")
def lockdown_status(db: Session = Depends(get_db), _=Depends(get_current_user)):
    state = db.query(SystemState).filter(SystemState.key == "lockdown").first()
    return {"active": state.value == "true" if state else False,
            "reason": state.detail if state else None,
            "since": state.updated_at.isoformat() if state else None}

@router.post("/lockdown")
def set_lockdown(body: LockdownRequest, db: Session = Depends(get_db), admin=Depends(require_admin)):
    state = db.query(SystemState).filter(SystemState.key == "lockdown").first()
    if not state:
        state = SystemState(key="lockdown", value="false")
        db.add(state)

    state.value      = "true" if body.active else "false"
    state.detail     = body.reason
    state.updated_by = admin.id
    state.updated_at = datetime.utcnow()

    # Lock / unlock all gates
    db.query(Gate).update({"locked": body.active})

    # Create alert
    alert = Alert(
        type="lockdown",
        description=f"{'🔴 EMERGENCY LOCKDOWN ACTIVATED' if body.active else '🟢 Lockdown lifted'} by {admin.full_name}. Reason: {body.reason}"
    )
    db.add(alert)
    audit(db, admin, "lockdown.activate" if body.active else "lockdown.lift", detail={"reason": body.reason})
    db.commit()
    return {"active": body.active, "message": "Lockdown activated" if body.active else "Lockdown lifted"}

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 3 — PHOTO UPLOAD
# ══════════════════════════════════════════════════════════════════════════════
PHOTO_DIR = "static/photos"
os.makedirs(PHOTO_DIR, exist_ok=True)

@router.post("/users/{user_id}/photo")
async def upload_photo(user_id: int, file: UploadFile = File(...),
                       db: Session = Depends(get_db), current=Depends(get_current_user)):
    # Students/staff can upload their own photo; admin can upload for anyone
    if current.role not in ("admin",) and current.id != user_id:
        raise HTTPException(403, "Can only upload your own photo")
    admin = current  # reuse audit variable
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(404, "User not found")
    if not file.content_type.startswith("image/"):
        raise HTTPException(400, "File must be an image")

    ext  = file.filename.rsplit(".", 1)[-1].lower()
    name = f"{user.qr_code}.{ext}"
    path = os.path.join(PHOTO_DIR, name)

    with open(path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    user.photo = f"/static/photos/{name}"
    audit(db, admin, "user.photo_upload", "user", user_id)
    db.commit()
    return {"photo_url": user.photo}

from fastapi.responses import FileResponse, Response
import hashlib, time

@router.get("/users/{user_id}/photo")
def get_photo(
    user_id: int,
    size: str = "full",   # full | thumb (80px) | small (160px)
    db: Session = Depends(get_db),
    _=Depends(get_current_user)
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user or not user.photo:
        raise HTTPException(404, "No photo")

    # Serve the right size
    base_path = user.photo.lstrip("/")
    size_path  = base_path.replace(".jpg", f"_{size}.jpg") if size != "full" else base_path

    if not os.path.exists(size_path):
        # Generate the requested size on-demand
        if os.path.exists(base_path) and size in ("thumb", "small"):
            from PIL import Image as PILImage
            import io as _io
            img  = PILImage.open(base_path).convert("RGB")
            dims = {"thumb": (80, 80), "small": (160, 160)}
            img.thumbnail(dims[size], PILImage.LANCZOS)
            img.save(size_path, format="JPEG", quality=80, optimize=True)
        elif os.path.exists(base_path):
            size_path = base_path
        else:
            raise HTTPException(404, "Photo file not found")

    # ETag based on file modification time
    mtime    = os.path.getmtime(size_path)
    etag     = hashlib.md5(f"{size_path}{mtime}".encode()).hexdigest()

    return FileResponse(
        size_path,
        media_type="image/jpeg",
        headers={
            "Cache-Control": "public, max-age=86400, stale-while-revalidate=3600",
            "ETag": etag,
            "Vary": "Accept-Encoding",
        }
    )
@router.delete("/users/{user_id}/photo")
def delete_photo(user_id: int, db: Session = Depends(get_db), current=Depends(get_current_user)):
    # Student can only delete their own; admin can delete any
    if current.role not in ("admin",) and current.id != user_id:
        raise HTTPException(403, "Can only delete your own photo")
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(404, "User not found")
    if user.photo:
        path = user.photo.lstrip("/")
        if os.path.exists(path): os.remove(path)
        user.photo = None
        audit(db, current, "user.photo_delete", "user", user_id)
        db.commit()
    return {"message": "Photo deleted"}



# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 4 — QR COMPROMISE / LOST CARD
# ══════════════════════════════════════════════════════════════════════════════
@router.post("/users/{user_id}/report-qr-lost")
def report_qr_lost(user_id: int, db: Session = Depends(get_db), current=Depends(get_current_user)):
    # Student can only report their own; admin can report any
    if current.role not in ("admin",) and current.id != user_id:
        raise HTTPException(403, "Can only report your own QR")

    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(404, "User not found")

    old_qr   = user.qr_code
    new_hash = generate_qr_hash(user.reg_number + secrets.token_hex(4))
    user.qr_code = new_hash

    # Remove old QR image
    old_path = f"static/qr_codes/{old_qr}.png"
    if os.path.exists(old_path): os.remove(old_path)
    generate_qr_image(new_hash, user.full_name)

    alert = Alert(user_id=user.id, type="qr_compromised",
                  description=f"{user.full_name} reported QR card lost/stolen. Old QR invalidated, new QR generated.")
    db.add(alert)
    audit(db, current, "user.qr_regenerate", "user", user_id, {"old_qr": old_qr, "new_qr": new_hash})
    db.commit()
    return {"message": "QR invalidated and regenerated", "new_qr": new_hash}

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 5 — GUARD STATISTICS
# ══════════════════════════════════════════════════════════════════════════════
@router.get("/stats/guards")
def guard_stats(db: Session = Depends(get_db), _=Depends(require_admin)):
    today = date.today()
    guards = db.query(User).filter(User.role == "guard").all()
    result = []
    for g in guards:
        approved_today = db.query(func.count(Entry.id)).filter(
            Entry.approved_by == g.id,
            func.date(Entry.scanned_at) == today
        ).scalar()
        approved_total = db.query(func.count(Entry.id)).filter(
            Entry.approved_by == g.id
        ).scalar()
        alerts_raised = db.query(func.count(Alert.id)).filter(
            func.date(Alert.timestamp) == today
        ).scalar()
        result.append({
            "id": g.id, "name": g.full_name, "status": g.status,
            "approved_today": approved_today,
            "approved_total": approved_total,
            "alerts_today": alerts_raised,
        })
    return result

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 7 — VISITOR PRE-REGISTRATION
# ══════════════════════════════════════════════════════════════════════════════
class PreRegRequest(BaseModel):
    visitor_name:  str
    visitor_email: Optional[str] = None
    visitor_phone: Optional[str] = None
    reason:        str
    expected_date: str   # YYYY-MM-DD
    expected_time: str   # HH:MM

@router.post("/visitors/pre-register")
def pre_register_visitor(body: PreRegRequest, db: Session = Depends(get_db), host=Depends(get_current_user)):
    token = secrets.token_urlsafe(32)
    expected_dt = datetime.strptime(f"{body.expected_date} {body.expected_time}", "%Y-%m-%d %H:%M")

    visitor = Visitor(
        name=body.visitor_name,
        email=body.visitor_email,
        phone=body.visitor_phone,
        reason=body.reason,
        host=host.full_name,
        host_user_id=host.id,
        temp_pass=secrets.token_hex(8),
        prereg_token=token,
        pre_registered=True,
        expected_exit=expected_dt,
        checked_out=False,
    )
    db.add(visitor); db.commit(); db.refresh(visitor)
    audit(db, host, "visitor.pre_register", "visitor", visitor.id,
          {"name": body.visitor_name, "date": body.expected_date})

    return {
        "visitor_id": visitor.id,
        "token": token,
        "checkin_link": f"/visitor-checkin/{token}",
        "message": f"Pre-registration created for {body.visitor_name}"
    }

@router.get("/visitors/checkin/{token}")
def visitor_checkin(token: str, db: Session = Depends(get_db)):
    visitor = db.query(Visitor).filter(Visitor.prereg_token == token).first()
    if not visitor: raise HTTPException(404, "Invalid or expired token")
    if visitor.checked_out: raise HTTPException(400, "Already checked out")
    return {
        "name": visitor.name, "reason": visitor.reason,
        "host": visitor.host, "expected_exit": visitor.expected_exit.isoformat() if visitor.expected_exit else None,
        "pre_registered": True
    }

@router.post("/visitors/checkin/{token}/confirm")
def confirm_visitor_checkin(token: str, db: Session = Depends(get_db), guard=Depends(require_guard)):
    visitor = db.query(Visitor).filter(Visitor.prereg_token == token).first()
    if not visitor: raise HTTPException(404, "Invalid token")
    visitor.time_in = datetime.utcnow()
    db.commit()
    return {"message": "Visitor checked in", "visitor_id": visitor.id}

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 8 — BULK USER IMPORT
# ══════════════════════════════════════════════════════════════════════════════
@router.post("/users/bulk-import")
async def bulk_import(file: UploadFile = File(...),
                      db: Session = Depends(get_db), admin=Depends(require_admin)):
    import openpyxl, io as _io
    content = await file.read()
    wb = openpyxl.load_workbook(_io.BytesIO(content))
    ws = wb.active

    created = []; skipped = []; errors = []
    # Auto-detect header row (scan first 10 rows for required columns)
    required = {"full_name", "email", "reg_number", "role", "password"}
    header_row_idx = None
    headers = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10), start=1):
        candidate = [str(c.value).strip().lower() if c.value else "" for c in row]
        if required.issubset(set(candidate)):
            header_row_idx = row_idx
            headers = candidate
            break

    if header_row_idx is None:
        raise HTTPException(400, f"Missing columns. Required: {', '.join(required)}. Headers must be in one of the first 10 rows.")

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not any(row): continue
        data = {headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)}
        if not data.get("full_name") or not data.get("reg_number"):
            skipped.append(data.get("reg_number", "unknown")); continue
        if db.query(User).filter(User.reg_number == data["reg_number"]).first():
            skipped.append(data["reg_number"]); continue
        if db.query(User).filter(User.email == data["email"]).first():
            skipped.append(data["reg_number"]); continue
        try:
            qr = generate_qr_hash(data["reg_number"])
            expiry = datetime.utcnow() + timedelta(days=180)  # 6-month QR validity
            u = User(full_name=data["full_name"], email=data["email"],
                     reg_number=data["reg_number"],
                     role=data.get("role","student").lower(),
                     password=hash_password(data.get("password","changeme123")),
                     qr_code=qr, qr_expires=expiry)
            db.add(u); db.flush()
            generate_qr_image(qr, data["full_name"])
            created.append(data["reg_number"])
        except Exception as e:
            errors.append({"reg": data.get("reg_number"), "error": str(e)})

    db.commit()
    audit(db, admin, "users.bulk_import", detail={"created": len(created), "skipped": len(skipped)})
    return {"created": len(created), "skipped": len(skipped), "errors": errors,
            "created_list": created[:20]}

@router.get("/users/bulk-template")
def download_template(_=Depends(require_admin)):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Users"
    headers = ["full_name","email","reg_number","role","password"]
    ws.append(headers)
    ws.append(["Amara Diallo","amara@bugema.ac.ug","23/BIT/BU/R/0001","student","changeme123"])
    ws.append(["Kevin Otieno","kotieno@bugema.ac.ug","LEC/2024/001","lecturer","changeme123"])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=campusguard_import_template.xlsx"})

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 9 — AUDIT LOG
# ══════════════════════════════════════════════════════════════════════════════
@router.get("/audit-logs")
def get_audit_logs(limit: int = 100, db: Session = Depends(get_db), _=Depends(require_admin)):
    logs = db.query(AuditLog).order_by(desc(AuditLog.timestamp)).limit(limit).all()
    return [
        {"id": l.id, "actor": l.actor_name, "action": l.action,
         "target_type": l.target_type, "target_id": l.target_id,
         "detail": l.detail, "timestamp": l.timestamp.strftime("%b %d, %Y – %I:%M %p")}
        for l in logs
    ]

# ══════════════════════════════════════════════════════════════════════════════
# FEATURE 12 — QR EXPIRY
# ══════════════════════════════════════════════════════════════════════════════
@router.post("/users/{user_id}/renew-qr")
def renew_qr(user_id: int, db: Session = Depends(get_db), admin=Depends(require_admin)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(404, "User not found")
    old_qr   = user.qr_code
    new_hash = generate_qr_hash(user.reg_number + secrets.token_hex(4))
    user.qr_code   = new_hash
    user.qr_expires = datetime.utcnow() + timedelta(days=180)
    old_path = f"static/qr_codes/{old_qr}.png"
    if os.path.exists(old_path): os.remove(old_path)
    generate_qr_image(new_hash, user.full_name)
    audit(db, admin, "user.qr_renew", "user", user_id)
    db.commit()
    return {"message": "QR renewed", "new_qr": new_hash, "expires": user.qr_expires.isoformat()}

@router.get("/users/expired-qr")
def list_expired_qr(db: Session = Depends(get_db), _=Depends(require_admin)):
    now = datetime.utcnow()
    expired = db.query(User).filter(User.qr_expires != None, User.qr_expires < now).all()
    expiring_soon = db.query(User).filter(
        User.qr_expires != None,
        User.qr_expires > now,
        User.qr_expires < now + timedelta(days=30)
    ).all()
    return {
        "expired": [{"id": u.id, "name": u.full_name, "reg": u.reg_number,
                     "expired_on": u.qr_expires.strftime("%b %d, %Y")} for u in expired],
        "expiring_soon": [{"id": u.id, "name": u.full_name, "reg": u.reg_number,
                           "expires": u.qr_expires.strftime("%b %d, %Y")} for u in expiring_soon],
    }

@router.post("/qr/check-expiry")
def auto_expire_qr(db: Session = Depends(get_db), _=Depends(require_admin)):
    """Run this periodically to auto-flag expired QRs."""
    now = datetime.utcnow()
    expired_users = db.query(User).filter(
        User.qr_expires != None, User.qr_expires < now, User.status == "active"
    ).all()
    flagged = []
    for u in expired_users:
        existing = db.query(Alert).filter(
            Alert.user_id == u.id, Alert.type == "qr_expired",
            Alert.resolved == False
        ).first()
        if not existing:
            db.add(Alert(user_id=u.id, type="qr_expired",
                         description=f"QR code for {u.full_name} ({u.reg_number}) expired on {u.qr_expires.strftime('%b %d, %Y')}. Renewal required."))
            flagged.append(u.full_name)
    db.commit()
    return {"flagged": len(flagged), "names": flagged}
